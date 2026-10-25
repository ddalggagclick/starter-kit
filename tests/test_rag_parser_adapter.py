from __future__ import annotations

import builtins
import importlib.util
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
from pathlib import Path
from types import SimpleNamespace
from typing import get_type_hints
from unittest.mock import AsyncMock, Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "rag"))

from docling.datamodel.base_models import FormatToExtensions, InputFormat
from docling_core.types.doc.base import BoundingBox
from docling_core.types.doc.document import DoclingDocument, ProvenanceItem
from docling_core.types.doc.items.table.table_data import TableCell, TableData
from docling_core.types.doc.labels import DocItemLabel
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from starter_rag import cli as rag_cli
from starter_rag.ingestion import parser
from starter_rag.ingestion.parser import (
    DocumentConversionError,
    DocumentElement,
    DocumentLocation,
    DocumentReadError,
    DocumentStructure,
    SourceProvenance,
    TableCellStructure,
    TableRowStructure,
    TableStructure,
    UnsupportedDocumentFormatError,
    extract_document_structure,
    parse_document,
)


def _cli_chunking_structure() -> DocumentStructure:
    table_location = DocumentLocation(
        status="unverified",
        page_numbers=(),
        provenance=(SourceProvenance(page_no=None, bbox=None, charspan=(0, 2)),),
    )
    table = DocumentElement(
        kind="table",
        text=None,
        original_text=None,
        self_ref="#/tables/0",
        level=None,
        location=table_location,
        table=TableStructure(
            row_count=1,
            column_count=1,
            rows=(TableRowStructure(row_index=0, cell_indexes=(0,)),),
            cells=(
                TableCellStructure(
                    text="안내",
                    start_row=0,
                    end_row=1,
                    start_column=0,
                    end_column=1,
                    row_span=1,
                    column_span=1,
                    column_header=True,
                    row_header=False,
                    row_section=False,
                    fillable=False,
                    self_ref="#/tables/0/cells/0",
                    bbox=None,
                ),
            ),
        ),
    )
    paragraph_location = DocumentLocation(status="not_applicable", page_numbers=(), provenance=())
    paragraphs = tuple(
        DocumentElement(
            kind="paragraph",
            text=text,
            original_text=text,
            self_ref=f"#/texts/{index}",
            level=None,
            location=paragraph_location,
            table=None,
        )
        for index, text in enumerate(("첫 문단", "둘째 문단"))
    )
    return DocumentStructure(
        source_format="xlsx", elements=(table, *paragraphs), unverified_items=()
    )


class DoclingParserAdapterTests(unittest.TestCase):
    @staticmethod
    def _write_pdf(path: Path, text: str) -> None:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
            ),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length "
            + str(len(content)).encode("ascii")
            + b" >>\nstream\n"
            + content
            + b"\nendstream",
        ]
        pdf = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, body in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{index} 0 obj\n".encode("ascii"))
            pdf.extend(body + b"\nendobj\n")
        xref_offset = len(pdf)
        pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
        pdf.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        pdf.extend(
            f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
        )
        path.write_bytes(pdf)

    def _assert_mocked_format(self, extension: str) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / f"synthetic{extension}"
            source.write_bytes(b"synthetic document")
            expected_document = object()
            converter = Mock()
            converter.convert.return_value = SimpleNamespace(
                status="success", document=expected_document
            )

            with patch.object(parser, "_get_converter", return_value=converter):
                document = parse_document(source)

        self.assertIs(document, expected_document)
        converter.convert.assert_called_once_with(source, raises_on_error=False)

    def test_설치된_Docling의_모든_확장자를_대소문자와_복합확장자까지_허용한다(self) -> None:
        extensions = sorted(
            {extension for values in FormatToExtensions.values() for extension in values}
        )
        self.assertEqual(set(FormatToExtensions), set(InputFormat))

        for extension in extensions:
            with self.subTest(extension=extension):
                self._assert_mocked_format(f".{extension}")

        for extension in (".XLSX", ".DCLG.XML", ".TAR.GZ"):
            with self.subTest(extension=extension):
                self._assert_mocked_format(extension)

    def test_문서변환기가_설치된_입력형식_전체를_지원한다(self) -> None:
        converter = parser._create_converter()

        self.assertEqual(set(converter.format_to_options), set(InputFormat))

    def test_실제_XLSX를_DoclingDocument로_변환한다(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic-guide.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet["A1"] = "Parser adapter check"
            worksheet["B1"] = "XLSX cell"
            workbook.save(source)

            document = parse_document(source)

        self.assertIsInstance(document, DoclingDocument)
        self.assertIn("Parser adapter check", document.export_to_markdown())

    def test_손상된_XLSX는_변환_오류로_구분한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "damaged.xlsx"
            source.write_bytes(b"not an Excel workbook")

            with self.assertRaises(DocumentConversionError):
                parse_document(source)

    def test_실제_PDF를_DoclingDocument로_변환한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic-guide.pdf"
            self._write_pdf(source, "PDF parser check")

            document = parse_document(source)

        self.assertIsInstance(document, DoclingDocument)
        self.assertIn("PDF parser check", document.export_to_markdown())

    def test_손상된_PDF는_변환_오류로_구분한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "damaged.pdf"
            source.write_bytes(b"%PDF-1.4\nthis is not a document")

            with self.assertRaises(DocumentConversionError):
                parse_document(source)

    def test_잘린_HTML은_Docling의_오류없이_복구되는_동작을_유지한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "truncated.html"
            source.write_text("<html><body><h1>Truncated HTML", encoding="utf-8")

            document = parse_document(source)

        self.assertIsInstance(document, DoclingDocument)
        self.assertIn("Truncated HTML", document.export_to_markdown())

    def test_합성_HTML을_실제_Docling으로_반복_변환한다(self) -> None:
        html = """<!doctype html>
<html><head><title>Synthetic guide</title></head>
<body><h1>Parser adapter check</h1><p>Generated only for this test.</p></body></html>
"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic-guide.html"
            source.write_text(html, encoding="utf-8")

            first = parse_document(source)
            second = parse_document(source)

        self.assertIsInstance(first, DoclingDocument)
        self.assertIs(get_type_hints(parse_document)["return"], DoclingDocument)
        self.assertIn("Parser adapter check", first.export_to_markdown())
        self.assertEqual(first.export_to_markdown(), second.export_to_markdown())
        structure = extract_document_structure(first, source)

        self.assertEqual(structure.source_format, "html")
        self.assertIn("Parser adapter check", [element.text for element in structure.elements])
        self.assertTrue(
            all(element.location.status == "not_applicable" for element in structure.elements)
        )

    def test_제목_절제목_본문을_원문_순서로_추출한다(self) -> None:
        document = DoclingDocument(name="synthetic structure")
        document.add_title("문서 제목")
        document.add_heading("첫 번째 절", level=2)
        document.add_text(label=DocItemLabel.PARAGRAPH, text="  원문 단락\n그대로  ")

        structure = extract_document_structure(document, "synthetic.pdf")

        self.assertEqual(
            [element.kind for element in structure.elements],
            ["document_title", "section_title", "paragraph"],
        )
        self.assertEqual(
            [element.text for element in structure.elements],
            ["문서 제목", "첫 번째 절", "  원문 단락\n그대로  "],
        )
        self.assertEqual(structure.elements[1].level, 2)
        self.assertTrue(all(element.self_ref for element in structure.elements))

    def test_목록_코드_체크박스_텍스트를_원문_순서로_보존한다(self) -> None:
        document = DoclingDocument(name="mixed labels")
        labels = (
            (DocItemLabel.LIST_ITEM, "A list entry with source wording."),
            (DocItemLabel.CODE, "example_value = 42"),
            (DocItemLabel.CHECKBOX_UNSELECTED, "□ Option not selected"),
            (DocItemLabel.CHECKBOX_SELECTED, "☑ Option selected"),
        )
        list_group = document.add_list_group()
        document.add_list_item(
            text=labels[0][1],
            parent=list_group,
            prov=ProvenanceItem(
                page_no=3,
                bbox=BoundingBox(l=0, t=0, r=100, b=20),
                charspan=(0, len(labels[0][1])),
            ),
        )
        for label, content in labels[1:]:
            document.add_text(label=label, text=content)

        structure = extract_document_structure(document, "synthetic.pdf")
        self.assertEqual(
            [element.text for element in structure.elements], [text for _, text in labels]
        )
        self.assertTrue(all(element.kind == "paragraph" for element in structure.elements))
        self.assertTrue(all(element.self_ref for element in structure.elements))
        self.assertEqual(structure.elements[0].location.page_numbers, (3,))
        self.assertEqual(
            structure.elements[0].location.provenance[0].charspan, (0, len(labels[0][1]))
        )

    def test_빈_표_구조를_보존하고_완전성_검토가_필요하다고_표시한다(self) -> None:
        document = DoclingDocument(name="empty table result")
        document.add_table(data=TableData(num_rows=0, num_cols=0, table_cells=[]))

        structure = extract_document_structure(document, "synthetic.pdf")

        self.assertEqual(len(structure.elements), 1)
        self.assertEqual(structure.elements[0].table.row_count, 0)
        self.assertEqual(structure.elements[0].table.column_count, 0)
        self.assertEqual(structure.elements[0].table.cells, ())
        self.assertEqual(
            {item.kind for item in structure.unverified_items},
            {"location", "table_completeness"},
        )

    def test_표의_행_열_셀과_병합_범위를_보존한다(self) -> None:
        document = DoclingDocument(name="synthetic table")
        document.add_table(
            data=TableData(
                num_rows=2,
                num_cols=2,
                table_cells=[
                    TableCell(
                        text="항목",
                        start_row_offset_idx=0,
                        end_row_offset_idx=1,
                        start_col_offset_idx=0,
                        end_col_offset_idx=2,
                        row_span=1,
                        col_span=2,
                        column_header=True,
                    ),
                    TableCell(
                        text="이름",
                        start_row_offset_idx=1,
                        end_row_offset_idx=2,
                        start_col_offset_idx=0,
                        end_col_offset_idx=1,
                        row_span=1,
                        col_span=1,
                    ),
                    TableCell(
                        text="홍길동",
                        start_row_offset_idx=1,
                        end_row_offset_idx=2,
                        start_col_offset_idx=1,
                        end_col_offset_idx=2,
                        row_span=1,
                        col_span=1,
                    ),
                ],
            )
        )

        structure = extract_document_structure(document, "synthetic.pdf")
        table_element = structure.elements[0]
        self.assertIsNotNone(table_element.table)
        table = table_element.table
        self.assertEqual((table.row_count, table.column_count), (2, 2))
        self.assertEqual([row.cell_indexes for row in table.rows], [(0,), (1, 2)])
        self.assertEqual([cell.text for cell in table.cells], ["항목", "이름", "홍길동"])
        self.assertEqual(table.cells[0].column_span, 2)
        self.assertEqual((table.cells[0].start_column, table.cells[0].end_column), (0, 2))
        self.assertTrue(table.cells[0].column_header)
        self.assertTrue(
            any(
                finding.kind == "table_completeness" and finding.self_ref == table_element.self_ref
                for finding in structure.unverified_items
            )
        )

    def test_페이지_provenance와_좌표를_검증된_위치로_보존한다(self) -> None:
        document = DoclingDocument(name="synthetic location")
        provenance = ProvenanceItem(
            page_no=4,
            bbox=BoundingBox(l=10, t=20, r=110, b=40),
            charspan=(2, 10),
        )
        document.add_heading("확인된 제목", level=1, prov=provenance)

        element = extract_document_structure(document, "synthetic.pdf").elements[0]

        self.assertEqual(element.location.status, "verified")
        self.assertEqual(element.location.page_numbers, (4,))
        self.assertEqual(element.location.provenance[0].charspan, (2, 10))
        self.assertEqual(
            element.location.provenance[0].bbox,
            parser.DocumentBoundingBox(10.0, 20.0, 110.0, 40.0, "TOPLEFT"),
        )

    def test_웹문서_HTML은_페이지_위치를_적용하지_않는다(self) -> None:
        document = DoclingDocument(name="synthetic HTML")
        document.add_text(
            label=DocItemLabel.PARAGRAPH,
            text="HTML 문단",
            prov=ProvenanceItem(
                page_no=1,
                bbox=BoundingBox(l=0, t=0, r=10, b=10),
                charspan=(0, 8),
            ),
        )

        element = extract_document_structure(document, "synthetic.html").elements[0]

        self.assertEqual(element.location.status, "not_applicable")
        self.assertEqual(element.location.page_numbers, ())
        self.assertEqual(element.location.provenance[0].page_no, 1)

    def test_위치_정보가_없으면_미검증으로_표시한다(self) -> None:
        document = DoclingDocument(name="synthetic missing provenance")
        document.add_text(label=DocItemLabel.PARAGRAPH, text="위치 정보 없음")

        structure = extract_document_structure(document, "synthetic.pdf")

        self.assertEqual(structure.elements[0].location.status, "unverified")
        self.assertEqual(structure.elements[0].location.page_numbers, ())
        self.assertEqual(structure.unverified_items[0].kind, "location")
        self.assertEqual(structure.unverified_items[0].self_ref, structure.elements[0].self_ref)

    def test_구조_추출은_반복해도_같은_결과를_반환한다(self) -> None:
        document = DoclingDocument(name="synthetic repeat")
        document.add_title("반복 검증")
        document.add_text(label=DocItemLabel.TEXT, text="같은 결과")

        first = extract_document_structure(document, "synthetic.docx")
        second = extract_document_structure(document, "synthetic.docx")

        self.assertEqual(first, second)

    def test_문서_PDF_경로를_변환기에_전달한다(self) -> None:
        self._assert_mocked_format(".pdf")

    def test_문서_DOCX_경로를_변환기에_전달한다(self) -> None:
        self._assert_mocked_format(".docx")

    def test_없는_파일은_별도_오류로_처리한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "missing.html"

            with self.assertRaises(FileNotFoundError):
                parse_document(source)

    def test_미지원_확장자는_별도_오류로_처리한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "notes.exe"
            source.write_text("synthetic", encoding="utf-8")

            with self.assertRaises(UnsupportedDocumentFormatError):
                parse_document(source)

    def test_읽을_수_없는_파일은_별도_오류로_처리한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "guide.html"
            source.write_text("synthetic", encoding="utf-8")

            with (
                patch.object(Path, "open", side_effect=PermissionError("denied")),
                patch.object(parser, "_get_converter") as get_converter,
                self.assertRaises(DocumentReadError),
            ):
                parse_document(source)

            get_converter.assert_not_called()

    def test_실패_상태의_Docling_결과를_오류로_처리한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "guide.html"
            source.write_text("synthetic", encoding="utf-8")
            converter = Mock()
            converter.convert.return_value = SimpleNamespace(
                status=SimpleNamespace(value="failure"), document=object()
            )

            with (
                patch.object(parser, "_get_converter", return_value=converter),
                self.assertRaises(DocumentConversionError),
            ):
                parse_document(source)

    def test_문서변환기_예외를_변환_오류로_감싼다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "guide.html"
            source.write_text("synthetic", encoding="utf-8")
            converter = Mock()
            converter.convert.side_effect = RuntimeError("synthetic conversion failure")

            with (
                patch.object(parser, "_get_converter", return_value=converter),
                self.assertRaises(DocumentConversionError) as context,
            ):
                parse_document(source)

        self.assertIsInstance(context.exception.__cause__, RuntimeError)

    def test_변환기_인스턴스를_재사용한다(self) -> None:
        converter = object()
        parser._get_converter.cache_clear()
        try:
            with patch.object(
                parser, "_create_converter", return_value=converter
            ) as create_converter:
                first = parser._get_converter()
                second = parser._get_converter()

            self.assertIs(first, second)
            create_converter.assert_called_once_with()
        finally:
            parser._get_converter.cache_clear()

    def test_모듈_가져오기는_Docling_변환이나_DB를_초기화하지_않는다(self) -> None:
        spec = importlib.util.spec_from_file_location(
            "isolated_rag_parser_import", Path(parser.__file__)
        )
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        isolated_parser = importlib.util.module_from_spec(spec)
        original_import = builtins.__import__

        def guarded_import(name, *args, **kwargs):
            if name == "docling" or name.startswith("docling."):
                raise AssertionError("The Docling converter must load only on request")
            if name == "sqlalchemy" or name.startswith("sqlalchemy."):
                raise AssertionError("Parser import must not load database code")
            return original_import(name, *args, **kwargs)

        sys.modules[spec.name] = isolated_parser
        try:
            with patch.object(builtins, "__import__", side_effect=guarded_import):
                spec.loader.exec_module(isolated_parser)
        finally:
            sys.modules.pop(spec.name, None)

        self.assertEqual(isolated_parser._get_converter.cache_info().currsize, 0)

    def test_명령행_파일_인자로_유효한_JSON을_출력한다(self) -> None:
        html = """<!doctype html><html><body>
<h1>CLI 구조 확인</h1><h2>제목</h2><p>본문</p>
<table><tr><th>이름</th></tr><tr><td>홍길동</td></tr></table>
</body></html>"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "cli-guide.html"
            source.write_text(html, encoding="utf-8")
            stdout, stderr = StringIO(), StringIO()

            with redirect_stdout(stdout), redirect_stderr(stderr):
                result = rag_cli.main(["parse", str(source)])

        self.assertEqual(result, 0)
        payload = json.loads(stdout.getvalue())
        self.assertEqual(len(stdout.getvalue().splitlines()), 1)
        self.assertEqual(payload["input_file"], str(source))
        self.assertTrue(payload["success"])
        self.assertGreater(payload["element_count"], 0)
        self.assertTrue(payload["elements"])
        self.assertTrue(
            all(
                element["location"]["status"] == "not_applicable" for element in payload["elements"]
            )
        )

    def test_파일_경로가_없거나_비어_있으면_사용법_오류를_낸다(self) -> None:
        for argv in ([], ["parse"]):
            stderr = StringIO()
            with redirect_stderr(stderr), self.assertRaises(SystemExit) as context:
                rag_cli.main(argv)
            self.assertEqual(context.exception.code, 2)
            self.assertIn("usage:", stderr.getvalue())

    def test_명령행에서_없는_파일은_종료_코드_2를_반환한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "missing.html"
            stdout, stderr = StringIO(), StringIO()

            with redirect_stdout(stdout), redirect_stderr(stderr):
                result = rag_cli.main(["parse", str(source)])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("문서 파일이 없습니다", stderr.getvalue())

    def test_명령행에서_미지원_형식은_종료_코드_2를_반환한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "notes.exe"
            source.write_text("synthetic", encoding="utf-8")
            stdout, stderr = StringIO(), StringIO()

            with redirect_stdout(stdout), redirect_stderr(stderr):
                result = rag_cli.main(["parse", str(source)])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("지원하지 않는 문서 형식", stderr.getvalue())

    def test_명령행에서_변환_오류는_종료_코드_1을_반환한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.object(
                rag_cli, "parse_document", side_effect=DocumentConversionError("변환 실패")
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["parse", "broken.pdf"])

        self.assertEqual(result, 1)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("문서 파싱에 실패했습니다", stderr.getvalue())

    def test_명령행에서_읽기_오류는_종료_코드_2를_반환한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.object(rag_cli, "parse_document", side_effect=DocumentReadError("읽기 실패")),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["parse", "unreadable.pdf"])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("읽기 실패", stderr.getvalue())

    def test_도움말에는_승인된_명령과_review를_표시한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            redirect_stdout(stdout),
            redirect_stderr(stderr),
            self.assertRaises(SystemExit) as context,
        ):
            rag_cli.main(["--help"])

        self.assertEqual(context.exception.code, 0)
        self.assertEqual(stderr.getvalue(), "")
        help_text = stdout.getvalue()
        command_line = next(line for line in help_text.splitlines() if "{" in line)
        for command in ("review", "config", "batch"):
            self.assertIn(command, command_line)
        for removed in ("validate-search", "source-v2", "evaluate-source-v2"):
            self.assertNotIn(removed, help_text)

    def test_이전_공개_명령과_옵션은_사용법_오류로_거부한다(self) -> None:
        for argv in (
            ["--file", "guide.pdf"],
            ["validate-search", "--query", "sample topic"],
            ["source-v2", "validate", "source.xlsx"],
            ["evaluate-source-v2"],
        ):
            stderr = StringIO()
            with redirect_stderr(stderr), self.assertRaises(SystemExit) as context:
                rag_cli.main(argv)
            self.assertEqual(context.exception.code, 2, argv)

    def test_파싱_명령은_위치_경로를_기존_파서_출력에_전달한다(self) -> None:
        with patch.object(rag_cli, "_parse_file", return_value=0) as parse_file:
            result = rag_cli.main(["parse", "manual.pdf"])

        self.assertEqual(result, 0)
        parse_file.assert_called_once_with("manual.pdf")

    def test_상태_명령이_구현과_환경설정_실행준비를_구분한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.dict(os.environ, {}, clear=True),
            patch.object(
                rag_cli,
                "_database_for_command",
                side_effect=AssertionError("status must not connect to the database"),
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["status"])

        self.assertEqual(result, 0)
        self.assertEqual(stderr.getvalue(), "")
        self.assertEqual(
            json.loads(stdout.getvalue()),
            {
                "package": "starter-rag",
                "status": "implemented",
                "searchImplemented": True,
                "ingestionImplemented": True,
                "databaseConfigured": False,
                "runtimeReadiness": "not-checked",
            },
        )

    def test_상태_명령은_유효한_로컬_DB_설정을_연결없이_확인한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.dict(
                os.environ,
                {
                    "RAG_DATABASE_URL": (
                        "postgresql+asyncpg://rag@/starter_rag_integration?host=%2Ftmp"
                    )
                },
                clear=True,
            ),
            patch.object(
                rag_cli,
                "_database_for_command",
                side_effect=AssertionError("status must not connect to the database"),
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["status"])

        self.assertEqual(result, 0)
        self.assertEqual(stderr.getvalue(), "")
        payload = json.loads(stdout.getvalue())
        self.assertTrue(payload["databaseConfigured"])
        self.assertEqual(payload["runtimeReadiness"], "not-checked")
        self.assertNotEqual(payload["status"], "ready")

    def test_상태_명령은_원격_DB_URL과_비밀번호를_출력하지_않는다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.dict(
                os.environ,
                {
                    "RAG_DATABASE_URL": (
                        "postgresql+asyncpg://user:status-secret@db.supabase.co:5432/postgres"
                    )
                },
                clear=True,
            ),
            patch.object(
                rag_cli,
                "_database_for_command",
                side_effect=AssertionError("status must not connect to the database"),
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["status"])

        self.assertEqual(result, 0)
        self.assertEqual(stderr.getvalue(), "")
        payload = json.loads(stdout.getvalue())
        self.assertFalse(payload["databaseConfigured"])
        self.assertEqual(payload["runtimeReadiness"], "not-checked")
        self.assertNotIn("status-secret", stdout.getvalue())

    def test_검색_명령은_검색어와_top_k를_DB_접속_전에_검증한다(self) -> None:
        for argv in (["search", "   "], ["search", "검색어", "--top-k", "11"]):
            stdout, stderr = StringIO(), StringIO()
            with (
                patch.object(
                    rag_cli,
                    "_database_for_command",
                    side_effect=AssertionError("잘못된 검색 요청은 DB에 접속하면 안 된다"),
                ),
                redirect_stdout(stdout),
                redirect_stderr(stderr),
            ):
                result = rag_cli.main(argv)

            self.assertEqual(result, 2, argv)
            self.assertEqual(stdout.getvalue(), "")
            self.assertIn("검색 요청", stderr.getvalue())

    def test_검색_명령은_위치_검색어와_기본_top_k를_기존_검색에_전달한다(self) -> None:
        with patch.object(rag_cli, "_search", return_value=0) as search:
            result = rag_cli.main(["search", "sample topic"])

        self.assertEqual(result, 0)
        search.assert_called_once_with("sample topic", 3, "published-guides", False)

    def test_검색_명령은_합성자료_허용과_collection_옵션을_전달한다(self) -> None:
        with patch.object(rag_cli, "_search", return_value=0) as search:
            result = rag_cli.main(
                [
                    "search",
                    "합성 시험 검색",
                    "--top-k",
                    "5",
                    "--collection-id",
                    "isolated-test",
                    "--allow-synthetic",
                ]
            )

        self.assertEqual(result, 0)
        search.assert_called_once_with("합성 시험 검색", 5, "isolated-test", True)

    def test_청킹은_명령줄_토크나이저_경로를_환경변수보다_우선한다(self) -> None:
        with (
            patch.dict(os.environ, {"RAG_TOKENIZER_PATH": "/env/tokenizer"}, clear=True),
            patch.object(rag_cli, "_chunk_file", return_value=0) as chunk_file,
        ):
            result = rag_cli.main(["chunk", "guide.pdf", "--tokenizer-path", "/cli/tokenizer"])

        self.assertEqual(result, 0)
        self.assertEqual(chunk_file.call_args.kwargs["tokenizer_path"], "/cli/tokenizer")

    def test_청킹은_RAG_TOKENIZER_PATH를_기본_토크나이저로_사용한다(self) -> None:
        with (
            patch.dict(os.environ, {"RAG_TOKENIZER_PATH": "/env/tokenizer"}, clear=True),
            patch.object(rag_cli, "_chunk_file", return_value=0) as chunk_file,
        ):
            result = rag_cli.main(["chunk", "guide.pdf"])

        self.assertEqual(result, 0)
        self.assertEqual(chunk_file.call_args.args, ("guide.pdf",))
        self.assertEqual(chunk_file.call_args.kwargs["tokenizer_path"], "/env/tokenizer")
        self.assertEqual(chunk_file.call_args.kwargs["target_tokens"], 512)

    def test_토크나이저_경로가_없으면_청킹_전에_입력_오류로_종료한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.dict(os.environ, {}, clear=True),
            patch.object(
                rag_cli,
                "_chunk_file",
                side_effect=AssertionError("설정 오류 뒤에는 청킹을 실행하면 안 된다"),
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["chunk", "guide.pdf"])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("RAG_TOKENIZER_PATH", stderr.getvalue())

    def test_청킹_도움말과_기존_status_명령이_정상_작동한다(self) -> None:
        help_stdout, help_stderr = StringIO(), StringIO()
        with (
            redirect_stdout(help_stdout),
            redirect_stderr(help_stderr),
            self.assertRaises(SystemExit) as context,
        ):
            rag_cli.main(["chunk", "--help"])

        self.assertEqual(context.exception.code, 0)
        self.assertIn("usage: starter-rag chunk", help_stdout.getvalue())
        self.assertEqual(help_stderr.getvalue(), "")

        status_stdout = StringIO()
        with redirect_stdout(status_stdout):
            result = rag_cli.main(["status"])
        self.assertEqual(result, 0)
        self.assertEqual(json.loads(status_stdout.getvalue())["package"], "starter-rag")

    def test_청킹_CLI가_모의_모델과_로컬_토크나이저로_검증용_JSON을_출력한다(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "chunk-guide.xlsx"
            source.write_bytes(b"test input")
            structure = _cli_chunking_structure()
            model = Mock(spec=BaseChatModel)

            async def generate(*_args, **_kwargs):
                print("모의 모델 진단")
                return AIMessage(content='{"groups":[["e000002","e000003"]]}')

            model.ainvoke = AsyncMock(side_effect=generate)

            def encode(text: str, *, add_special_tokens: bool) -> list[str]:
                print("로컬 토크나이저 진단")
                self.assertFalse(add_special_tokens)
                return text.split()

            def load_local_tokenizer(_path: str) -> SimpleNamespace:
                print("토크나이저 로드 진단")
                return SimpleNamespace(encode=encode)

            stdout, stderr = StringIO(), StringIO()
            with (
                patch.object(rag_cli, "parse_document", return_value=object()),
                patch.object(rag_cli, "extract_document_structure", return_value=structure),
                patch.object(rag_cli, "_load_tokenizer", side_effect=load_local_tokenizer),
                patch(
                    "starter_rag.ingestion.service.create_chunking_model_from_environment",
                    return_value=model,
                ),
                redirect_stdout(stdout),
                redirect_stderr(stderr),
            ):
                result = rag_cli.main(
                    [
                        "chunk",
                        str(source),
                        "--tokenizer-path",
                        "local-tokenizer",
                    ]
                )

        self.assertEqual(result, 0, stderr.getvalue())
        self.assertTrue(stdout.getvalue(), repr(stderr.getvalue()))
        payload = json.loads(stdout.getvalue())
        self.assertEqual(len(stdout.getvalue().splitlines()), 1)
        self.assertTrue(payload["success"])
        self.assertEqual(payload["chunk_count"], 2)
        self.assertEqual(payload["tokenizer"], "local-tokenizer")
        table_chunk = next(chunk for chunk in payload["chunks"] if chunk["kind"] == "table_row")
        self.assertEqual(table_chunk["table_row"]["row_index"], 0)
        self.assertEqual(table_chunk["table_row"]["cells"][0]["cell"]["text"], "안내")
        self.assertEqual(
            table_chunk["locations"][0]["provenance"][0]["charspan"],
            [0, 2],
        )
        self.assertIn("모의 모델 진단", stderr.getvalue())
        self.assertIn("로컬 토크나이저 진단", stderr.getvalue())
        self.assertIn("토크나이저 로드 진단", stderr.getvalue())
        model.ainvoke.assert_awaited_once()

    def test_청킹_CLI_잘못된_토크나이저_경로는_비밀값_없이_종료_코드_2를_반환한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.object(rag_cli, "parse_document", return_value=object()),
            patch.object(
                rag_cli,
                "extract_document_structure",
                return_value=DocumentStructure(
                    source_format="pdf", elements=(), unverified_items=()
                ),
            ),
            patch.object(rag_cli, "_load_tokenizer", side_effect=ValueError("tokenizer-secret")),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["chunk", "guide.pdf", "--tokenizer-path", "/missing/tokenizer"])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("토크나이저", stderr.getvalue())
        self.assertNotIn("tokenizer-secret", stderr.getvalue())

    def test_청킹_CLI_잘못된_공급자_설정은_종료_코드_2를_반환한다(self) -> None:
        stdout, stderr = StringIO(), StringIO()
        with (
            patch.dict(os.environ, {"RAG_CHUNKING_PROVIDER": "invalid-secret"}, clear=True),
            patch.object(rag_cli, "parse_document", return_value=object()),
            patch.object(
                rag_cli,
                "extract_document_structure",
                return_value=DocumentStructure(
                    source_format="pdf", elements=(), unverified_items=()
                ),
            ),
            patch.object(
                rag_cli,
                "_load_tokenizer",
                return_value=SimpleNamespace(encode=lambda text, **_: text.split()),
            ),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            result = rag_cli.main(["chunk", "guide.pdf", "--tokenizer-path", "local-tokenizer"])

        self.assertEqual(result, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertTrue(stderr.getvalue().strip())
        self.assertNotIn("invalid-secret", stderr.getvalue())

    def test_명령행_모듈_가져오기는_파싱이나_DB_초기화를_하지_않는다(self) -> None:
        spec = importlib.util.spec_from_file_location(
            "isolated_rag_cli_import", Path(rag_cli.__file__)
        )
        self.assertIsNotNone(spec)
        self.assertIsNotNone(spec.loader)
        isolated_cli = importlib.util.module_from_spec(spec)
        original_import = builtins.__import__

        def guarded_import(name, *args, **kwargs):
            if name == "docling" or name.startswith("docling."):
                raise AssertionError("CLI import must not load Docling")
            if name == "sqlalchemy" or name.startswith("sqlalchemy."):
                raise AssertionError("CLI import must not load database code")
            return original_import(name, *args, **kwargs)

        sys.modules[spec.name] = isolated_cli
        try:
            with (
                patch.object(builtins, "__import__", side_effect=guarded_import),
                patch.object(parser, "_get_converter") as get_converter,
            ):
                spec.loader.exec_module(isolated_cli)
                get_converter.assert_not_called()
        finally:
            sys.modules.pop(spec.name, None)


if __name__ == "__main__":
    unittest.main()
